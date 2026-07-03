#!/usr/bin/env python3
"""
slides_to_pptx.py
ChatGPTで生成したスライド画像（PNG/JPG）を PowerPoint ファイルに自動変換する

使い方：
  python3 slides_to_pptx.py                      # slides/ フォルダを使う（デフォルト）
  python3 slides_to_pptx.py ./my_slides           # フォルダを指定
  python3 slides_to_pptx.py ./my_slides out.pptx  # 出力ファイル名も指定
"""

import os
import glob
import sys
from pathlib import Path


def create_presentation(slides_dir: str, output_path: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(10)       # 16:9 横幅
    prs.slide_height = Inches(5.625)   # 16:9 高さ

    # PNG / JPG をファイル名順に取得
    patterns = ["*.png", "*.jpg", "*.jpeg", "*.PNG", "*.JPG", "*.JPEG"]
    image_files = []
    for pat in patterns:
        image_files.extend(glob.glob(os.path.join(slides_dir, pat)))
    image_files = sorted(set(image_files))  # 重複除去してソート

    if not image_files:
        print(f"エラー：{slides_dir} に画像が見つかりません。")
        print("slide_01.png ～ slide_10.png のようなファイルを slides/ フォルダに入れてください。")
        sys.exit(1)

    print(f"画像を {len(image_files)} 枚検出しました。")

    blank_layout = prs.slide_layouts[6]  # 空白レイアウト

    for i, img_path in enumerate(image_files, 1):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            left=0,
            top=0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        print(f"  [{i:02d}] {os.path.basename(img_path)}")

    # 出力先ディレクトリを作成
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    prs.save(output_path)
    print(f"\n完了 → {output_path}  ({len(image_files)} 枚)")
    print("このファイルを Google スライドにアップロードするか、PowerPoint で開いてください。")


def main():
    slides_dir = sys.argv[1] if len(sys.argv) > 1 else "./slides"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "./output/training_material.pptx"

    if not os.path.isdir(slides_dir):
        print(f"エラー：フォルダが見つかりません → {slides_dir}")
        print("使い方：python3 slides_to_pptx.py [スライドフォルダ] [出力ファイル名]")
        sys.exit(1)

    print(f"スライドフォルダ：{slides_dir}")
    print(f"出力先          ：{output_path}")
    print("-" * 40)
    create_presentation(slides_dir, output_path)


if __name__ == "__main__":
    main()
